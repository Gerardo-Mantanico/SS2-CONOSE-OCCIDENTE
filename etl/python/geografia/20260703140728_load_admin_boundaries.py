#!/usr/bin/env python3

"""
ETL de límites administrativos (OCHA/HDX Common Operational Dataset).

Reproducible y trazable:
  1. Lee un archivo Excel con la estructura estándar COD (hojas *_admin0/1/2).
  2. GENERA un script SQL determinista en scripts/generated/geografia/
  3. EJECUTA el SQL generado en la base.

Genérico por país: descubre las hojas por sufijo (admin0/admin1/admin2), así
sirve para Guatemala, otro país u otra versión sin cambiar el código —basta
cambiar el archivo de entrada (misma estructura) y la fuente.

El SQL generado:
  - resuelve el padre por 'pcode' (no por id),
  - Si el pcode ya existe hace un UPDATE por lo que se puede re-correr sin problema,
  - registra la carga en meta.carga (con SHA-256 del archivo) y su cobertura en meta.cobertura_carga,

Uso:
    python 20260703140000_load_admin_boundaries.py
    python 20260703140000_load_admin_boundaries.py --input data/geografia/hnd_admin.xlsx --fuente HND_ADMIN_BOUNDARIES
    python 20260703140000_load_admin_boundaries.py --sql-only # genera sin ejecutar
"""

from __future__ import annotations

import argparse
import hashlib
import logging
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

# Ubicar raíz del repo e importar config
def _repo_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / ".env").exists() or (parent / "scripts").is_dir():
            return parent
    return Path(__file__).resolve().parents[3]


REPO_ROOT = _repo_root()
for _p in Path(__file__).resolve().parents:
    if (_p / "config.py").exists():
        sys.path.insert(0, str(_p))
        break

from config import get_connection

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-7s  %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("etl.geografia")

# --- Valores por defecto -----------------------------
# Archivo fuente versionado en el repo para que el flujo sea reproducible.
DEFAULT_INPUT = REPO_ROOT / "data" / "geografia" / "2026_07_gtm_admin_boundaries.xlsx"
# Debe coincidir con un 'codigo' registrado en data/fuentes.csv.
DEFAULT_FUENTE = "2026_07_gtm_admin_boundaries"

OUTPUT_DIR = REPO_ROOT / "scripts" / "generated" / "geografia"


# --- Utilidades ---------------------------------------------------------------
def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()

def _sql_str(s) -> str:
    return "NULL" if s in (None, "") else "'" + str(s).replace("'", "''") + "'"

def _sql_num(x) -> str:
    return "NULL" if x in (None, "") else repr(float(x))

def _excel_date(v) -> str:
    """Formatea una fecha como DATE 'YYYY-MM-DD'. Acepta datetime/date (lo que
    suele devolver openpyxl) o un serial de Excel; NULL si no se puede."""
    if v in (None, ""):
        return "NULL"
    if isinstance(v, (datetime, date)):
        return f"DATE '{v:%Y-%m-%d}'"
    try:
        d = datetime(1899, 12, 30) + timedelta(days=int(float(v)))
        return f"DATE '{d:%Y-%m-%d}'"
    except (TypeError, ValueError):
        return "NULL"


def _read_sheets(path: Path) -> dict:
    """Lee las hojas admin0/1/2 por sufijo (genérico por país)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    hojas = {}
    for nombre in wb.sheetnames:
        low = nombre.lower()
        for nivel in ("admin0", "admin1", "admin2"):
            if low.endswith(nivel):
                hojas[nivel] = nombre
    faltan = {"admin0", "admin1", "admin2"} - hojas.keys()
    if faltan:
        raise ValueError(f"Al archivo le faltan hojas: {sorted(faltan)}. "
                         f"Hojas encontradas: {wb.sheetnames}")

    def filas(sheet):
        ws = wb[sheet]
        data = list(ws.iter_rows(values_only=True))
        idx = {h: i for i, h in enumerate(data[0])}
        return idx, [r for r in data[1:] if r and any(c not in (None, "") for c in r)]

    return {nivel: filas(hojas[nivel]) for nivel in ("admin0", "admin1", "admin2")}


# --- Generación del SQL -------------------------------------------------------
def construir_sql(input_path: Path, fuente_codigo: str) -> tuple[str, dict]:
    hojas = _read_sheets(input_path)
    i0, d0 = hojas["admin0"]
    i1, d1 = hojas["admin1"]
    i2, d2 = hojas["admin2"]

    p = d0[0]
    iso3 = (p[i0["iso3"]] or "").upper()
    version = p[i0.get("version", -1)] if "version" in i0 else None
    valid_on = _excel_date(p[i0["valid_on"]]) if "valid_on" in i0 else "NULL"

    pais = (_sql_str(p[i0["adm0_name"]]), _sql_str(p[i0["adm0_pcode"]]),
            _sql_str(p[i0["iso2"]]), _sql_str(p[i0["iso3"]]),
            _sql_num(p[i0["area_sqkm"]]), _sql_num(p[i0["center_lat"]]),
            _sql_num(p[i0["center_lon"]]))

    dep = [f"        ({_sql_str(r[i1['adm1_name']])},{_sql_str(r[i1['adm1_pcode']])},"
           f"{_sql_num(r[i1['area_sqkm']])},{_sql_num(r[i1['center_lat']])},"
           f"{_sql_num(r[i1['center_lon']])})"
           for r in sorted(d1, key=lambda r: r[i1["adm1_pcode"]])]

    mun = [f"        ({_sql_str(r[i2['adm2_name']])},{_sql_str(r[i2['adm2_pcode']])},"
           f"{_sql_str(r[i2['adm1_pcode']])},{_sql_num(r[i2['area_sqkm']])},"
           f"{_sql_num(r[i2['center_lat']])},{_sql_num(r[i2['center_lon']])})"
           for r in sorted(d2, key=lambda r: r[i2["adm2_pcode"]])]

    total = 1 + len(dep) + len(mun)
    sha = _sha256(input_path)
    NL = ",\n"

    sql = f"""-- GENERADO AUTOMÁTICAMENTE
-- Generador: etl/python/geografia/{Path(__file__).name}
-- Fuente:    {input_path.name}
-- SHA-256:   {sha}
-- País:      {iso3}   Versión: {version}
-- Filas:     país=1, departamentos={len(dep)}, municipios={len(mun)} (total={total})
--
-- El padre se resuelve por 'pcode'. Idempotente: ON CONFLICT (pcode) DO UPDATE.
-- SUPUESTOS de esquema:
--   pais(nombre,pcode,iso2,iso3,area_km2,centro_lat,centro_lon,carga_id)
--   departamento(nombre,pcode,area_km2,centro_lat,centro_lon,pais_id,carga_id)
--   municipio(nombre,pcode,area_km2,centro_lat,centro_lon,departamento_id,carga_id)
--   pcode UNIQUE en las tres (migración V20260703120000__add_campos_hdx_cod.sql).

BEGIN;

WITH carga AS (
    INSERT INTO meta.carga
        (fuente_id, estado_carga_id, nombre_script, archivo_fuente, hash_archivo,
         ejecutado_por, finalizado_en, filas_procesadas, filas_insertadas)
    SELECT f.id, e.id, 'geografia/{Path(__file__).name}', {_sql_str(input_path.name)},
           '{sha}', current_user, now(), {total}, {total}
    FROM meta.fuente f
    JOIN meta.estado_carga e ON e.nombre = 'exitosa'
    WHERE f.codigo = {_sql_str(fuente_codigo)}
    RETURNING id
),
cobertura AS (
    INSERT INTO meta.cobertura_carga
        (carga_id, schema_destino, tabla_destino, periodo_inicio)
    SELECT (SELECT id FROM carga), 'geografia', t, {valid_on}
    FROM (VALUES ('pais'), ('departamento'), ('municipio')) AS x(t)
    RETURNING 1
),
pais AS (
    INSERT INTO geografia.pais
        (nombre, pcode, iso2, iso3, area_km2, centro_lat, centro_lon, carga_id)
    SELECT {pais[0]}, {pais[1]}, {pais[2]}, {pais[3]}, {pais[4]}, {pais[5]}, {pais[6]},
           (SELECT id FROM carga)
    ON CONFLICT (pcode) DO UPDATE SET
        nombre = EXCLUDED.nombre, iso2 = EXCLUDED.iso2, iso3 = EXCLUDED.iso3,
        area_km2 = EXCLUDED.area_km2, centro_lat = EXCLUDED.centro_lat,
        centro_lon = EXCLUDED.centro_lon, carga_id = EXCLUDED.carga_id
    RETURNING id
),
dep AS (
    INSERT INTO geografia.departamento
        (nombre, pcode, area_km2, centro_lat, centro_lon, pais_id, carga_id)
    SELECT d.nombre, d.pcode, d.area_km2, d.centro_lat, d.centro_lon,
           (SELECT id FROM pais), (SELECT id FROM carga)
    FROM (VALUES
{NL.join(dep)}
    ) AS d(nombre, pcode, area_km2, centro_lat, centro_lon)
    ON CONFLICT (pcode) DO UPDATE SET
        nombre = EXCLUDED.nombre, area_km2 = EXCLUDED.area_km2,
        centro_lat = EXCLUDED.centro_lat, centro_lon = EXCLUDED.centro_lon,
        pais_id = EXCLUDED.pais_id, carga_id = EXCLUDED.carga_id
    RETURNING id, pcode
)
INSERT INTO geografia.municipio
    (nombre, pcode, area_km2, centro_lat, centro_lon, departamento_id, carga_id)
SELECT m.nombre, m.pcode, m.area_km2, m.centro_lat, m.centro_lon,
       dep.id, (SELECT id FROM carga)
FROM (VALUES
{NL.join(mun)}
    ) AS m(nombre, pcode, adm1_pcode, area_km2, centro_lat, centro_lon)
JOIN dep ON dep.pcode = m.adm1_pcode
ON CONFLICT (pcode) DO UPDATE SET
    nombre = EXCLUDED.nombre, area_km2 = EXCLUDED.area_km2,
    centro_lat = EXCLUDED.centro_lat, centro_lon = EXCLUDED.centro_lon,
    departamento_id = EXCLUDED.departamento_id, carga_id = EXCLUDED.carga_id;

COMMIT;
"""
    meta = {"iso3": iso3 or "XXX", "version": version or "v00",
            "total": total, "dep": len(dep), "mun": len(mun), "sha": sha}
    return sql, meta


def ejecutar_sql(sql: str) -> None:
    conn = get_connection()
    try:
        conn.autocommit = True  # el propio SQL controla la transacción (BEGIN/COMMIT)
        with conn.cursor() as cur:
            cur.execute(sql)
    finally:
        conn.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Carga de límites administrativos (COD).")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT,
                        help=f"Archivo Excel de entrada (def: {DEFAULT_INPUT})")
    parser.add_argument("--fuente", default=DEFAULT_FUENTE,
                        help=f"Código de fuente en meta.fuente (def: {DEFAULT_FUENTE})")
    parser.add_argument("--sql-only", action="store_true",
                        help="Genera el SQL pero no lo ejecuta.")
    args = parser.parse_args()

    if not args.input.exists():
        raise FileNotFoundError(
            f"No se encontró el archivo de entrada: {args.input}\n"
            "Coloca el Excel ahí (y versiónalo) o pásalo con --input."
        )

    log.info(f"Leyendo {args.input.name} (fuente '{args.fuente}')...")
    sql, meta = construir_sql(args.input, args.fuente)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / f"load_geografia_{meta['iso3'].lower()}_{meta['version']}.sql"
    out_path.write_text(sql, encoding="utf-8")
    log.info(f"SQL generado: {out_path.relative_to(REPO_ROOT)}  "
             f"(país=1, dep={meta['dep']}, muni={meta['mun']}, sha={meta['sha'][:12]}…)")

    if args.sql_only:
        log.info("--sql-only: no se ejecuta.")
        return

    log.info("Ejecutando contra la base...")
    ejecutar_sql(sql)
    log.info(f"Carga completada: {meta['total']} filas.")


if __name__ == "__main__":
    main()